from __future__ import annotations

import csv
import hashlib
import io
import json
import re
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from orchestrator.config import get_settings


TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".log", ".py", ".js", ".ts", ".tsx", ".jsx", ".css", ".html", ".htm", ".json", ".csv"}
ALLOWED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx"}


class _TextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())


def safe_name(name: str) -> str:
    clean = re.sub(r"[^\w.()\- ]+", "_", Path(name).name, flags=re.UNICODE).strip(" .")
    return clean[:220] or "material.txt"


def extract_text(name: str, data: bytes) -> tuple[str, dict[str, object]]:
    ext = Path(name).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise ValueError("Unsupported file type. Use TXT, Markdown, JSON, CSV, HTML, PDF or DOCX.")
    meta: dict[str, object] = {"extension": ext}
    if ext == ".pdf":
        reader = PdfReader(io.BytesIO(data))
        meta["pages"] = len(reader.pages)
        text = "\n\n".join(page.extract_text() or "" for page in reader.pages)
    elif ext == ".docx":
        document = Document(io.BytesIO(data))
        meta["paragraphs"] = len(document.paragraphs)
        text = "\n".join(item.text for item in document.paragraphs)
    else:
        decoded = data.decode("utf-8-sig", errors="replace")
        if ext == ".json":
            try:
                decoded = json.dumps(json.loads(decoded), ensure_ascii=False, indent=2)
            except json.JSONDecodeError as exc:
                raise ValueError(f"Invalid JSON: {exc.msg}") from exc
        elif ext == ".csv":
            rows = list(csv.reader(io.StringIO(decoded)))
            decoded = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
            meta["rows"] = len(rows)
        elif ext in {".html", ".htm"}:
            parser = _TextExtractor()
            parser.feed(decoded)
            decoded = "\n".join(parser.parts)
        text = decoded
    text = re.sub(r"\x00", "", text).strip()
    if not text:
        raise ValueError("The file contains no readable text.")
    meta["characters"] = len(text)
    return text[:500_000], meta


def persist_file(asset_id: str, name: str, data: bytes) -> tuple[str, str]:
    settings = get_settings()
    directory = Path(settings.file_storage_dir).resolve()
    directory.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(data).hexdigest()
    destination = directory / f"{asset_id}-{safe_name(name)}"
    destination.write_bytes(data)
    return str(destination), digest


def remove_file(path: str) -> None:
    target = Path(path).resolve()
    root = Path(get_settings().file_storage_dir).resolve()
    if root == target or root not in target.parents:
        return
    target.unlink(missing_ok=True)
